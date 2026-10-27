"""Fixed A4 workbook, with original data kept in a separate traceability table."""
from collections import defaultdict
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.pagebreak import Break
from domain import export_rows, calculate

THIN = Side(style='thin', color='243645')
YELLOW = PatternFill('solid', fgColor='FFF9B0')

def text(value):
    # Prevent spreadsheet formula injection in source-derived text.
    value = str(value or '')
    return "'" + value if value.startswith(('=', '+', '-', '@')) else value

def workbook_bytes(project, draft=False):
    rows = export_rows(project['rows'], draft)
    files = {f['id']: f['name'] for f in project.get('files', [])}
    wb = Workbook()
    wb.remove(wb.active)
    groups = defaultdict(list)
    for row in sorted(rows, key=lambda r: (r['building'], r['floor'], r['system'], r['room'], r['direction'], r['device'])):
        groups[(row['building'], row['floor'], row['system'])].append(row)
    # One source row per device or undivided group; never invent per-device distributions.
    for (building, floor, system), group in groups.items():
        paired = []
        for row in group:
            match = next((x for x in paired if x['room'] == row['room'] and x['name'] == row['room_name']
                          and x['scope'] == row['scope'] and x['file'] == row['source_file']
                          and x['page'] == row['source_page'] and row['direction'] not in x['points']), None)
            if match is None:
                match = {'room': row['room'], 'name': row['room_name'], 'scope': row['scope'],
                         'file': row['source_file'], 'page': row['source_page'], 'points': {}}
                paired.append(match)
            match['points'][row['direction']] = row
        for start in range(0, len(paired), 14):
            page = paired[start:start+14]
            ws = wb.create_sheet(f'Protokoll {len(wb.worksheets)+1:02d}')
            ws.sheet_view.showGridLines = False
            widths = [10, 25, 11, 11, 11, 11, 11, 11, 21]
            for i, width in enumerate(widths, 1):
                ws.column_dimensions[chr(64+i)].width = width
            def box(area, value, fill=False, bold=False, size=10):
                ws.merge_cells(area)
                for line in ws[area]:
                    for cell in line:
                        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
                        cell.alignment = Alignment(vertical='center', wrap_text=True)
                        if fill: cell.fill = YELLOW
                cell = ws[area.split(':')[0]]
                cell.value = text(value)
                cell.font = Font(name='Arial', size=size, bold=bold, color='172F3F')
            box('A1:D2', 'Luftmätningsprotokoll' + (' – UTKAST' if draft else ''), bold=True, size=16)
            box('E1:G1', 'Plan / byggnad')
            box('E2:G2', f'{floor} / {building}', True, True)
            box('H1:I1', 'System')
            box('H2:I2', system, True, True)
            box('A4:D4', 'Fastighet / uppdrag')
            box('A5:D5', project['name'], True)
            box('E4:G4', 'Mättekniker')
            box('E5:G5', project.get('technician', ''), True)
            box('H4:I4', 'Mätdatum')
            box('H5:I5', project.get('date', ''), True)
            box('A7:D7', 'Flödesenhet: l/s')
            box('E7:I7', 'Andel av projekterat = uppmätt / projekterat × 100')
            headers = ['Rum', 'Benämning / omfattning', 'Proj.\nTilluft', 'Uppmätt\nTilluft', '% av proj.\nTilluft',
                       'Proj.\nFrånluft', 'Uppmätt\nFrånluft', '% av proj.\nFrånluft', 'Källa']
            for c, value in enumerate(headers, 1):
                ws.cell(9, c, value)
            ws.row_dimensions[9].height = 34
            for n, item in enumerate(page, 10):
                ws.cell(n, 1, text(item['room']))
                ws.cell(n, 2, text(item['name'] + '\n' + item['scope']))
                citations = []
                for direction, cols in [('T', (3, 4, 5)), ('F', (6, 7, 8))]:
                    row = item['points'].get(direction)
                    if not row: continue
                    values = calculate(row)
                    ws.cell(n, cols[0], values['design_ls'])
                    ws.cell(n, cols[1], values['measured_ls'])
                    design_col, measured_col = chr(64+cols[0]), chr(64+cols[1])
                    ws.cell(n, cols[2], f'=IF(AND(ISNUMBER({design_col}{n}),{design_col}{n}>0,ISNUMBER({measured_col}{n})),{measured_col}{n}/{design_col}{n},"")')
                    ws.cell(n, cols[2]).number_format = '0.0%'
                    for c in cols[:2]: ws.cell(n, c).number_format = '0.0'
                    citation = f'{files.get(row["source_file"], "Manuell uppgift")}, s. {row["source_page"]}'
                    if citation not in citations: citations.append(citation)
                ws.cell(n, 9, text('\n'.join(citations)))
                ws.row_dimensions[n].height = 38
            end = 9 + len(page)
            for line in ws.iter_rows(min_row=9, max_row=end, max_col=9):
                for cell in line:
                    cell.font = Font(name='Arial', size=10, bold=cell.row == 9)
                    cell.alignment = Alignment(vertical='center', wrap_text=True, horizontal='left' if cell.column in (1,2,9) else 'center')
                    cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            box(f'A{end+2}:I{end+2}', 'Tom procentsiffra: projekterat jämförelsevärde saknas eller är noll. m³/h ÷ 3,6 = l/s.')
            ws.row_dimensions[end+2].height = 28
            box(f'A{end+4}:D{end+4}', 'Namnteckning')
            box(f'A{end+5}:D{end+6}', '')
            for n in [1,2,4,5,7]: ws.row_dimensions[n].height = 23
            ws.print_area = f'A1:I{end+7}'
            ws.page_setup.orientation = 'portrait'
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 1
            ws.page_margins.left = ws.page_margins.right = 0.25
            ws.page_margins.top = ws.page_margins.bottom = 0.3
            ws.print_options.horizontalCentered = True
            ws.oddFooter.center.text = 'Sida &P av &N'
    source = wb.create_sheet('Mätunderlag')
    columns = ['Byggnad','Plan','System','Rum','Benämning','Luftslag','Don','Omfattning','Proj. original','Enhet','Uppmätt original','Enhet','Granskad','Källfil','Sida']
    source.append(columns)
    for row in rows:
        source.append([text(row[k]) for k in ('building','floor','system','room','room_name','direction','device','scope')]
                      + [row['design'],row['design_unit'],row['measured'],row['measured_unit'],'Ja' if row['reviewed'] else 'Nej',text(files.get(row['source_file'], 'Manuell')),row['source_page']])
    source.freeze_panes = 'E2'
    source.auto_filter.ref = source.dimensions
    for col in source.columns:
        source.column_dimensions[col[0].column_letter].width = 19
    for cell in source[1]: cell.font = Font(bold=True)
    result = BytesIO()
    wb.save(result)
    return result.getvalue()
