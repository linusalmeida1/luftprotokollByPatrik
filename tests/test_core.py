import base64
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from domain import normalize_row, calculate, export_rows
from exporter import workbook_bytes
import app
import ai

def row(**changes):
    return normalize_row(dict(room='101',room_name='Kontor',building='A',floor='Plan 1',system='LB01',scope='Rumstotal',direction='T',design=150,design_unit='m³/h',measured=43,measured_unit='l/s',reviewed=True,**changes)) if not changes else normalize_row({**dict(room='101',room_name='Kontor',building='A',floor='Plan 1',system='LB01',scope='Rumstotal',direction='T',design=150,design_unit='m³/h',measured=43,measured_unit='l/s',reviewed=True),**changes})

class Calculations(unittest.TestCase):
    def test_conversion_and_percentage(self):
        self.assertAlmostEqual(calculate(row())['percentage'],1.032)
        self.assertAlmostEqual(calculate(row(design=40,design_unit='l/s',measured=144,measured_unit='m³/h'))['percentage'],1)
    def test_zero_is_measured(self):
        r=row(measured=0)
        self.assertEqual(calculate(r)['percentage'],0)
        self.assertEqual(len(export_rows([r])),1)
    def test_missing_not_zero(self):
        self.assertIsNone(calculate(row(design=None))['percentage'])
        self.assertIsNone(calculate(row(design=0))['percentage'])
        with self.assertRaises(ValueError):export_rows([row(measured=None)])
    def test_unknown_unit_blocks_final(self):
        with self.assertRaises(ValueError):export_rows([row(design_unit='')])
        with self.assertRaises(ValueError):export_rows([row(reviewed=False)])
        self.assertEqual(len(export_rows([row(reviewed=False)],True)),1)
    def test_exclusion_and_scope(self):
        self.assertEqual(len(export_rows([row(),row(measured=None,excluded=True)])),1)
        with self.assertRaises(ValueError):export_rows([row(scope='')])
    def test_number_validation(self):
        self.assertEqual(row(measured='12,5')['measured'],12.5)
        for bad in (-1,'NaN','Infinity','12 l/s',True):
            with self.assertRaises(ValueError):row(measured=bad)
    def test_excel_layout_and_formula_injection(self):
        p={'name':'=BAD()','rows':[row(room='001',room_name='=HYPERLINK("bad")')],'files':[]}
        w=load_workbook(BytesIO(workbook_bytes(p)))
        s=w.worksheets[0]
        self.assertEqual(str(s.page_setup.paperSize),'9')
        self.assertEqual(s.page_setup.orientation,'portrait')
        self.assertEqual(s.page_setup.fitToWidth,1)
        self.assertEqual(s['A10'].value,'001')
        self.assertEqual(s['B10'].data_type,'s')
        self.assertEqual(s['A5'].data_type,'s')
        self.assertEqual(s['E10'].data_type,'f')
        self.assertAlmostEqual(s['C10'].value,150/3.6)
    def test_pagination_and_mixed_groups(self):
        p={'name':'Test','rows':[row(room=str(i)) for i in range(35)],'files':[]}
        w=load_workbook(BytesIO(workbook_bytes(p)))
        self.assertGreaterEqual(len(w.worksheets),3)
        self.assertEqual(w['Mätunderlag'].max_row,36)

class Storage(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.old=app.DATA;app.DATA=Path(self.temp.name)
    def tearDown(self):app.DATA=self.old;self.temp.cleanup()
    def test_revision_conflict(self):
        p=app.new_project('Test');app.mutate(p['id'],lambda p:p.update(name='Nytt'),1)
        with self.assertRaises(ValueError):app.mutate(p['id'],lambda p:p.update(name='Förlorat'),1)
        self.assertEqual(app.read_project(p['id'])['name'],'Nytt')
    def test_upload_validation(self):
        raw,mime,name=app.validate_file('../../x.pdf',base64.b64encode(b'%PDF-1.7 test').decode())
        self.assertEqual(name,'x.pdf');self.assertEqual(mime,'application/pdf')
        with self.assertRaises(ValueError):app.validate_file('x.png',base64.b64encode(b'<html>bad').decode())
    def test_ai_import_is_unreviewed_and_idempotent(self):
        p=app.new_project('Test');fid='file-test';(app.DATA/'files').mkdir();(app.DATA/'files'/fid).write_bytes(b'%PDF-1.7 test')
        app.mutate(p['id'],lambda p:p['files'].append({'id':fid,'name':'x.pdf','mime':'application/pdf','analyzed':False}))
        app.JOBS['test']={'status':'running'}
        with patch.object(ai,'extract',return_value=[row()]):app.run_job('test',p['id'],fid)
        loaded=app.read_project(p['id']);self.assertFalse(loaded['rows'][0]['reviewed']);self.assertEqual(len(loaded['rows']),1)
        with patch.object(ai,'extract',return_value=[row()]):app.run_job('test',p['id'],fid)
        self.assertEqual(len(app.read_project(p['id'])['rows']),1)

if __name__=='__main__':unittest.main()
